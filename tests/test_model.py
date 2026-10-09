
"""Independent validation tests for the annuity model."""

from dataclasses import replace
from math import isclose, isfinite
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from src.assumptions import (
    Assumptions,
    base_assumptions,
    load_interest_rates,
)
from src.annuity_model import (
    project_annuity,
    summarize_projection,
)
from src.generate_excel import generate_workbook


class TestAssumptions(unittest.TestCase):

    def test_valid_base_assumptions(self):
        assumptions = base_assumptions()

        self.assertTrue(assumptions.validate())
        self.assertEqual(assumptions.years, 10)
        self.assertEqual(assumptions.initial_policies, 100)

    def test_invalid_assumptions_are_rejected(self):
        invalid_changes = [
            {"initial_premium": -100},
            {"initial_policies": 0},
            {"years": 0},
            {"years": 2.5},
            {"investment_rate": -1.0},
            {"credited_rate": -0.01},
            {"withdrawal_rate": 1.5},
            {"surrender_rate": 1.1},
            {"surrender_charge_rate": -0.1},
            {"expense_rate": float("nan")},
        ]

        for changes in invalid_changes:
            with self.subTest(changes=changes):
                assumptions = replace(
                    Assumptions(),
                    **changes,
                )

                with self.assertRaises(ValueError):
                    assumptions.validate()

    def test_interest_rate_source(self):
        observations = load_interest_rates()

        self.assertGreater(len(observations), 0)

        dates = [item["date"] for item in observations]

        self.assertEqual(dates, sorted(dates))
        self.assertEqual(len(dates), len(set(dates)))

        self.assertEqual(
            observations[-1]["series"],
            "DGS10",
        )

        self.assertTrue(
            isclose(
                base_assumptions().investment_rate,
                observations[-1]["rate_percent"] / 100.0,
            )
        )


class TestFinancialCalculations(unittest.TestCase):

    def test_known_interest_and_investment_example(self):
        # Independent expected calculation:
        # Premium: 1000
        # Investment income: 1000 * 5% = 50
        # Credited interest: 1000 * 3% = 30
        # Insurer surplus: 1050 - 1030 = 20

        assumptions = Assumptions(
            initial_premium=1000,
            initial_policies=10,
            years=1,
            investment_rate=0.05,
            credited_rate=0.03,
            withdrawal_rate=0,
            surrender_rate=0,
            surrender_charge_rate=0,
            expense_rate=0,
        )

        row = project_annuity(assumptions)[0]

        self.assertAlmostEqual(
            row["investment_income"], 50
        )
        self.assertAlmostEqual(
            row["credited_interest"], 30
        )
        self.assertAlmostEqual(
            row["closing_assets"], 1050
        )
        self.assertAlmostEqual(
            row["closing_balance"], 1030
        )
        self.assertAlmostEqual(
            row["closing_surplus"], 20
        )

    def test_partial_withdrawal_accounting(self):
        # 10% of 1000 = 100 withdrawn.
        # Assets and policyholder balances both fall to 900.

        assumptions = Assumptions(
            initial_premium=1000,
            years=1,
            investment_rate=0,
            credited_rate=0,
            withdrawal_rate=0.10,
            surrender_rate=0,
            surrender_charge_rate=0,
            expense_rate=0,
        )

        row = project_annuity(assumptions)[0]

        self.assertAlmostEqual(
            row["partial_withdrawals"], 100
        )
        self.assertAlmostEqual(
            row["closing_balance"], 900
        )
        self.assertAlmostEqual(
            row["closing_assets"], 900
        )
        self.assertAlmostEqual(
            row["closing_surplus"], 0
        )

    def test_full_surrender_accounting(self):
        # 20% of 1000 is surrendered: 200.
        # Charge: 10% of 200 = 20.
        # Cash paid: 180.
        # Remaining liability: 800.
        # Remaining insurer assets: 820.

        assumptions = Assumptions(
            initial_premium=1000,
            initial_policies=100,
            years=1,
            investment_rate=0,
            credited_rate=0,
            withdrawal_rate=0,
            surrender_rate=0.20,
            surrender_charge_rate=0.10,
            expense_rate=0,
        )

        row = project_annuity(assumptions)[0]

        self.assertAlmostEqual(
            row["gross_surrenders"], 200
        )
        self.assertAlmostEqual(
            row["surrender_charges"], 20
        )
        self.assertAlmostEqual(
            row["surrender_payouts"], 180
        )
        self.assertAlmostEqual(
            row["closing_balance"], 800
        )
        self.assertAlmostEqual(
            row["closing_assets"], 820
        )
        self.assertAlmostEqual(
            row["closing_surplus"], 20
        )
        self.assertAlmostEqual(
            row["closing_policies"], 80
        )

    def test_combined_known_example(self):
        # Independently calculated example:
        # Opening premium: 1000
        # Investment income: 60
        # Credited interest: 40
        # Post-credit balance: 1040
        # Partial withdrawals: 104
        # Remaining balance: 936
        # Gross surrenders: 187.20
        # Surrender charge: 9.36
        # Surrender cash payment: 177.84
        # Operating expenses: 10
        # Closing assets: 768.16
        # Closing policyholder balance: 748.80
        # Surplus: 19.36

        assumptions = Assumptions(
            initial_premium=1000,
            years=1,
            investment_rate=0.06,
            credited_rate=0.04,
            withdrawal_rate=0.10,
            surrender_rate=0.20,
            surrender_charge_rate=0.05,
            expense_rate=0.01,
        )

        row = project_annuity(assumptions)[0]

        expected_values = {
            "investment_income": 60,
            "credited_interest": 40,
            "post_credit_balance": 1040,
            "partial_withdrawals": 104,
            "gross_surrenders": 187.20,
            "surrender_charges": 9.36,
            "surrender_payouts": 177.84,
            "operating_expenses": 10,
            "cash_inflows": 60,
            "cash_outflows": 291.84,
            "net_cash_flow": -231.84,
            "closing_balance": 748.80,
            "closing_assets": 768.16,
            "closing_surplus": 19.36,
        }

        for key, expected in expected_values.items():
            with self.subTest(measure=key):
                self.assertAlmostEqual(
                    row[key], expected, places=7
                )

    def test_withdrawals_occur_after_interest(self):
        assumptions = Assumptions(
            initial_premium=1000,
            years=1,
            investment_rate=0,
            credited_rate=0.04,
            withdrawal_rate=0.10,
            surrender_rate=0,
            surrender_charge_rate=0,
            expense_rate=0,
        )

        row = project_annuity(assumptions)[0]

        # Withdrawal must be 10% of 1040, not 1000.
        self.assertAlmostEqual(
            row["partial_withdrawals"], 104
        )

    def test_balance_and_cash_flow_reconciliation(self):
        records = project_annuity(base_assumptions())

        self.assertEqual(len(records), 10)

        for row in records:

            expected_interest = (
                row["opening_balance"]
                * base_assumptions().credited_rate
            )

            expected_investment = (
                row["opening_assets"]
                * base_assumptions().investment_rate
            )

            self.assertTrue(
                isclose(
                    row["credited_interest"],
                    expected_interest,
                    abs_tol=1e-7,
                )
            )

            self.assertTrue(
                isclose(
                    row["investment_income"],
                    expected_investment,
                    abs_tol=1e-7,
                )
            )

            self.assertAlmostEqual(
                row["closing_balance"],
                row["opening_balance"]
                + row["credited_interest"]
                - row["partial_withdrawals"]
                - row["gross_surrenders"],
                places=7,
            )

            self.assertAlmostEqual(
                row["gross_surrenders"],
                row["surrender_payouts"]
                + row["surrender_charges"],
                places=7,
            )

            self.assertAlmostEqual(
                row["cash_outflows"],
                row["partial_withdrawals"]
                + row["surrender_payouts"]
                + row["operating_expenses"],
                places=7,
            )

            self.assertAlmostEqual(
                row["closing_assets"],
                row["opening_assets"]
                + row["cash_inflows"]
                - row["cash_outflows"],
                places=7,
            )

            self.assertAlmostEqual(
                row["annual_profit_change"],
                row["investment_income"]
                - row["credited_interest"]
                + row["surrender_charges"]
                - row["operating_expenses"],
                places=7,
            )

            for value in row.values():
                self.assertTrue(isfinite(value))

    def test_year_to_year_continuity(self):
        records = project_annuity(base_assumptions())

        for previous, current in zip(
            records[:-1], records[1:]
        ):
            self.assertAlmostEqual(
                current["opening_balance"],
                previous["closing_balance"],
                places=7,
            )

            self.assertAlmostEqual(
                current["opening_assets"],
                previous["closing_assets"],
                places=7,
            )

            self.assertAlmostEqual(
                current["opening_policies"],
                previous["closing_policies"],
                places=7,
            )

    def test_cumulative_surplus_reconciliation(self):
        records = project_annuity(base_assumptions())
        summary = summarize_projection(records)

        self.assertAlmostEqual(
            summary["ending_insurer_surplus"],
            summary["total_change_in_surplus"],
            places=7,
        )

        self.assertAlmostEqual(
            summary["ending_insurer_assets"],
            summary["ending_account_balance"]
            + summary["ending_insurer_surplus"],
            places=7,
        )

    def test_higher_investment_return_increases_surplus(self):
        base = base_assumptions()

        lower = replace(
            base,
            investment_rate=base.investment_rate - 0.01,
        )

        higher = replace(
            base,
            investment_rate=base.investment_rate + 0.01,
        )

        lower_result = summarize_projection(
            project_annuity(lower)
        )

        higher_result = summarize_projection(
            project_annuity(higher)
        )

        self.assertGreater(
            higher_result["ending_insurer_surplus"],
            lower_result["ending_insurer_surplus"],
        )


class TestExcelWorkbook(unittest.TestCase):

    def test_workbook_contains_real_formulas(self):
        with TemporaryDirectory() as folder:

            filename = Path(folder) / "test_annuity.xlsx"

            generate_workbook(filename)

            workbook = load_workbook(
                filename,
                data_only=False,
            )

            self.assertEqual(
                workbook.sheetnames,
                [
                    "Guide",
                    "Assumptions",
                    "Rates",
                    "Projection",
                    "Summary",
                ],
            )

            self.assertEqual(
                workbook["Assumptions"]["B4"].value,
                "=Rates!C8/100",
            )

            self.assertEqual(
                workbook["Projection"]["E5"].value,
                "=B5*Assumptions!$B$5",
            )

            self.assertEqual(
                workbook["Projection"]["S5"].value,
                "=R5-Q5",
            )

            self.assertEqual(
                workbook["Projection"]["A14"].value,
                10,
            )

            self.assertEqual(
                workbook["Summary"]["B19"].value,
                "=B7-B8+B9-B10-B3",
            )

            self.assertEqual(
                len(workbook["Summary"]._charts),
                2,
            )

            workbook.close()


if __name__ == "__main__":
    unittest.main()
